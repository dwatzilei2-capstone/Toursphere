"""Read actual application exports; verify types, IDs, totals, pages and content."""
from pathlib import Path
import json, csv, math
from datetime import datetime
import openpyxl
import pdfplumber
import subprocess

root = Path(__file__).resolve().parents[1]
folder = root / 'tmp/reports-validation'
checks = 0
def check(condition, label):
    global checks
    assert condition, label
    checks += 1

for manifest_file in folder.glob('*.json'):
    m = json.loads(manifest_file.read_text(encoding='utf-8'))
    name = manifest_file.stem
    wb = openpyxl.load_workbook(folder / (name + '.xlsx'))
    check(wb.sheetnames[0] == 'Summary', name + ' summary first')
    summary = wb['Summary']
    for i, metric in enumerate(m['metrics'], 9):
        cell = summary.cell(i, 2)
        if metric['value'] is not None:
            check(isinstance(cell.value, (int, float)), name + ' metric numeric')
            check(math.isclose(cell.value, metric['value']), name + ' metric exact')
    for index, (key, s) in enumerate(m['sections'].items(), 1):
        sheet = wb.worksheets[index]
        check(sheet.freeze_panes == 'A7', name + ' freeze')
        if s['count']: check(sheet.auto_filter.ref is not None, name + ' filter')
        rows = [json.loads(line) for line in (root / 'storage/private/reports' / m['token'] / (key + '.jsonl')).read_text(encoding='utf-8').splitlines()]
        for i, row in enumerate(rows, 7):
            for j, (field, col) in enumerate(s['columns'].items(), 1):
                value = row.get(field)
                cell = sheet.cell(i, j)
                if value is None or value == '': continue
                if col['type'] in ('integer', 'decimal', 'money', 'percent'):
                    check(isinstance(cell.value, (int, float)), name + ' numeric ' + field)
                    check(math.isclose(cell.value, value), name + ' exact numeric ' + field)
                    if col['type'] == 'money': check('₱' in cell.number_format, name + ' peso format')
                elif col['type'] in ('date', 'datetime'):
                    check(isinstance(cell.value, datetime), name + ' date ' + field)
                    check(cell.value.strftime('%Y-%m-%d') == value[:10], name + ' date exact')
                else:
                    check(cell.value == str(value), name + ' text exact ' + field + ' ' + str(cell.value))
                    if col['type'] == 'id': check(cell.data_type == 's' and cell.font.bold, name + ' ID typed and emphasized')
        for field, total in (s['totals'] or {}).items():
            colindex = list(s['columns']).index(field) + 1
            check(math.isclose(sheet.cell(7 + s['count'], colindex).value, total), name + ' grand total')
    with (folder / (name + '.csv')).open(encoding='utf-8-sig', newline='') as f:
        data = list(csv.reader(f))
    s = m['sections']['records']
    check(data[0] == [c['label'] for c in s['columns'].values()], name + ' CSV headers')
    check(len(data)-1 == s['count'], name + ' CSV count')
    ids = [json.loads(line)[next(iter(s['columns']))] for line in (root / 'storage/private/reports' / m['token'] / 'records.jsonl').read_text(encoding='utf-8').splitlines()]
    check([r[0] for r in data[1:]] == [str(v) for v in ids], name + ' CSV IDs')
    records = [json.loads(line) for line in (root / 'storage/private/reports' / m['token'] / 'records.jsonl').read_text(encoding='utf-8').splitlines()]
    for record, values in zip(records, data[1:]):
        for (field, col), actual in zip(s['columns'].items(), values):
            expected = record.get(field)
            if expected is None: check(actual == '', name + ' CSV missing data'); continue
            if col['type'] in ('integer', 'decimal', 'money', 'percent'):
                check(math.isclose(float(actual), expected), name + ' CSV numeric matches snapshot')
            else:
                import re
                safe = "'" + str(expected) if re.match(r'^\s*[=+@-]', str(expected)) else str(expected)
                check(actual == safe, name + ' CSV escaped text matches snapshot')
    pdf = pdfplumber.open(folder / (name + '.pdf'))
    text = '\n'.join(page.extract_text() for page in pdf.pages)
    check(m['title'] in text, name + ' PDF title')
    for value in ids: check(str(value) in text, name + ' PDF ID')
    for i, page in enumerate(pdf.pages):
        check(f'Page {i+1} of {len(pdf.pages)}' in page.extract_text(), name + ' page numbering')
        for block in page.extract_words():
            check(block['x0'] >= 20 and block['x1'] <= page.width-20 and block['top'] >= 10 and block['bottom'] <= page.height-15, name + ' PDF bounds')
    # Render representative summary and detail pages for visual inspection.
    if name in ('fleet', 'fuel', 'trips'):
        subprocess.run(['pdftoppm', '-f', '1', '-l', '2', '-scale-to', '1500', '-png', str(folder / (name + '.pdf')), str(folder / name)], check=True, capture_output=True)
pdf = pdfplumber.open(folder / 'multipage.pdf')
check(len(pdf.pages) > 4, 'Fixture genuinely spans multiple table pages')
for i, page in enumerate(pdf.pages):
    text = page.extract_text()
    check(f'Page {i+1} of {len(pdf.pages)}' in text, 'Multipage numbering')
    if i: check('Trip ID' in text, 'Repeated table header')
text = '\n'.join(page.extract_text() for page in pdf.pages)
for i in range(1, 36): check('TRIP-' + str(i).zfill(5) in text, 'Multipage ID retained')
print(f'{checks} actual workbook, PDF and CSV validation checks passed.')
